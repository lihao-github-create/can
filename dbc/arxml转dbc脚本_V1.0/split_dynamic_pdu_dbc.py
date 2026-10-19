#!/usr/bin/env python3
from __future__ import annotations

import argparse
import copy
import math
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

import canmatrix.convert
import canmatrix.formats
from canmatrix.canmatrix import ArbitrationId, CanMatrix, Frame, Signal
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


HEADER_SIGNAL_NAMES = ("Header_ID", "Header_DLC")
AUTOSAR_I_SIGNAL_I_PDU = "I-SIGNAL-I-PDU"
MAPPING_HEADERS = [
    "SourceFrameName",
    "SourceFrameIdHex",
    "SourceFrameType",
    "ContainerName",
    "SourceGroupName",
    "NewMsgName",
    "NewMsgIdHex",
    "NewMsgType",
    "NewMsgDlc",
    "SignalCount",
    "ShiftBits",
]


@dataclass
class MappingRow:
    source_frame_name: str
    source_frame_id_hex: str
    source_frame_type: str
    container_name: str
    source_group_name: str
    new_msg_name: str
    new_msg_id_hex: str
    new_msg_type: str
    new_msg_dlc: int
    signal_count: int
    shift_bits: int

    def to_row(self) -> list[str | int]:
        return [
            self.source_frame_name,
            self.source_frame_id_hex,
            self.source_frame_type,
            self.container_name,
            self.source_group_name,
            self.new_msg_name,
            self.new_msg_id_hex,
            self.new_msg_type,
            self.new_msg_dlc,
            self.signal_count,
            self.shift_bits,
        ]


@dataclass
class SplitStats:
    dynamic_frame_count: int = 0
    preserved_grouped_frame_count: int = 0
    created_split_count: int = 0
    duplicate_split_count: int = 0
    skipped_dynamic_group_count: int = 0


@dataclass
class ExportValidationIssue:
    frame_name: str
    signal_name: str
    factor: str
    offset: str
    minimum: str
    maximum: str
    error_type: str
    error_message: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Split dynamic-PDU container messages in DBC files into standalone "
            "messages, or convert an ARXML to DBC first and then split the "
            "dynamic PDUs in one command."
        )
    )
    parser.add_argument(
        "inputs",
        nargs="*",
        help=(
            "Split-only mode: one or more input DBC file paths or wildcard patterns. "
            "If omitted and --arxml is given, the script runs in convert+split mode."
        ),
    )
    parser.add_argument(
        "--arxml",
        required=True,
        help=(
            "AUTOSAR ARXML file. "
            "Used to look up exact PDU lengths in all modes, and used as the "
            "direct conversion source in convert+split mode."
        ),
    )
    parser.add_argument(
        "--output-dir",
        default=".",
        help="Directory for generated files. Default: current directory.",
    )
    parser.add_argument(
        "--output-prefix",
        default=None,
        help=(
            "Convert+split mode only. Prefix for generated raw/split DBC files. "
            "Example: PPD_HEV_V04 -> PPD_HEV_V04_ADCANFD.dbc. "
            "Default: ARXML file stem."
        ),
    )
    parser.add_argument(
        "--suffix",
        default="_split",
        help="Suffix appended to each generated split DBC stem. Default: _split",
    )
    parser.add_argument(
        "--id-strategy",
        choices=("header-id-extended", "duplicate-container-id"),
        default="header-id-extended",
        help=(
            "How to assign message IDs to dynamic split PDUs. "
            "'header-id-extended' keeps the output DBC ID-unique by using the "
            "contained PDU Header_ID as an extended-frame ID. "
            "'duplicate-container-id' preserves the source container CAN ID and "
            "may create duplicate BO_ IDs."
        ),
    )
    parser.add_argument(
        "--group-frame-id-mode",
        choices=("keep-original", "keep-original-extended"),
        default="keep-original",
        help=(
            "Legacy option for assigning IDs to split messages created from "
            "SIG_GROUP_ frames that do not contain Header_ID/Header_DLC. "
            "Grouped frames without Header_ID/Header_DLC are now preserved "
            "without splitting. "
            "'keep-original' preserves the original frame ID and frame type. "
            "'keep-original-extended' preserves the original numeric ID but forces "
            "the generated message to use an extended frame."
        ),
    )
    parser.add_argument(
        "--mapping-xlsx",
        dest="mapping_xlsx",
        action="store_true",
        default=True,
        help="Generate one mapping workbook per split DBC. Default: enabled.",
    )
    parser.add_argument(
        "--no-mapping-xlsx",
        dest="mapping_xlsx",
        action="store_false",
        help="Only generate DBC output and skip the mapping workbook.",
    )
    parser.add_argument(
        "--keep-converted-dbc",
        dest="keep_converted_dbc",
        action="store_true",
        default=True,
        help=(
            "Convert+split mode: also keep the raw DBC files produced directly "
            "from ARXML. Default: enabled."
        ),
    )
    parser.add_argument(
        "--no-keep-converted-dbc",
        dest="keep_converted_dbc",
        action="store_false",
        help="Convert+split mode: do not save the intermediate raw DBC files.",
    )
    parser.add_argument(
        "--dbc-export-encoding",
        default="iso-8859-1",
        help="DBC export encoding passed to canmatrix. Default: iso-8859-1",
    )
    parser.add_argument(
        "--dbc-export-comment-encoding",
        default="iso-8859-1",
        help="DBC comment export encoding passed to canmatrix. Default: iso-8859-1",
    )
    parser.add_argument(
        "--zero-factor-policy",
        choices=("error", "set-one"),
        default="error",
        help=(
            "How to handle signals whose scaling factor is 0 before DBC export. "
            "'error' stops with a detailed report. "
            "'set-one' rewrites factor 0 to 1 in memory so export can continue. "
            "Default: error"
        ),
    )
    return parser.parse_args()


def validate_args(args: argparse.Namespace) -> str:
    if args.inputs:
        return "split-only"
    return "convert-and-split"


def warn_if_risky_options(args: argparse.Namespace) -> None:
    if args.id_strategy == "duplicate-container-id":
        print(
            "Warning: --id-strategy duplicate-container-id preserves the source "
            "container CAN ID for all split dynamic PDUs. This can create duplicate "
            "BO_ IDs in the output DBC and may prevent canmatrix from re-reading the "
            "generated file reliably.",
            file=sys.stderr,
        )


def expand_dbc_inputs(patterns: Iterable[str]) -> list[Path]:
    resolved: list[Path] = []
    seen: set[Path] = set()
    for raw_pattern in patterns:
        pattern = Path(raw_pattern)
        matches = [pattern] if pattern.exists() else list(Path().glob(raw_pattern))
        if not matches:
            raise FileNotFoundError(f"DBC input not found: {raw_pattern}")
        for match in matches:
            full = match.resolve()
            if full.suffix.lower() != ".dbc":
                raise ValueError(f"Expected a .dbc file, got: {match}")
            if full not in seen:
                seen.add(full)
                resolved.append(full)
    return resolved


def parse_pdu_lengths(arxml_path: Path) -> dict[str, int]:
    lengths: dict[str, int] = {}
    namespace = ""
    for event, elem in ET.iterparse(arxml_path, events=("start", "end")):
        if event == "start" and not namespace and elem.tag.startswith("{"):
            namespace = elem.tag.split("}", 1)[0] + "}"
        if event != "end" or elem.tag != f"{namespace}{AUTOSAR_I_SIGNAL_I_PDU}":
            continue

        short_name = elem.findtext(f"{namespace}SHORT-NAME")
        length_text = elem.findtext(f"{namespace}LENGTH")
        if short_name and length_text:
            lengths[short_name] = int(length_text)
        elem.clear()
    if not lengths:
        raise ValueError(f"No {AUTOSAR_I_SIGNAL_I_PDU} definitions found in {arxml_path}")
    return lengths


def load_can_matrix(dbc_path: Path) -> tuple[str, CanMatrix]:
    cluster = canmatrix.formats.loadp(str(dbc_path))
    if not cluster:
        raise ValueError(f"Unable to load DBC: {dbc_path}")
    if len(cluster) != 1:
        raise ValueError(
            f"Expected exactly one CAN matrix in {dbc_path}, got {len(cluster)}"
        )
    raw_cluster_name, matrix = next(iter(cluster.items()))
    cluster_name = str(raw_cluster_name) or dbc_path.stem
    return cluster_name, matrix


def load_can_cluster_from_arxml(arxml_path: Path) -> dict[str, CanMatrix]:
    cluster = canmatrix.formats.loadp(str(arxml_path))
    if not cluster:
        raise ValueError(f"Unable to load ARXML: {arxml_path}")
    matrices = {str(name): matrix for name, matrix in cluster.items()}
    if not matrices:
        raise ValueError(f"No CAN clusters found in ARXML: {arxml_path}")
    for matrix in matrices.values():
        convert_pdu_container_frames(matrix)
        apply_preferred_signal_names(matrix)
    return matrices


def convert_pdu_container_frames(matrix: CanMatrix) -> None:
    frame_pdu_container_list = [
        frame
        for frame in matrix.frames
        if frame.is_pdu_container
    ]
    for frame in frame_pdu_container_list:
        new_frame = canmatrix.convert.convert_pdu_container_to_multiplexed(frame)
        matrix.del_frame(frame)
        matrix.add_frame(new_frame)


def preferred_signal_name(signal: Signal) -> str:
    sys_signal_name = str(signal.attributes.get("SysSignalName", "")).strip()
    return sys_signal_name or signal.name


def apply_preferred_signal_names(matrix: CanMatrix) -> None:
    for frame in matrix.frames:
        renamed: list[tuple[Signal, str]] = []
        seen_names: dict[str, str] = {}
        for signal in frame.signals:
            new_name = preferred_signal_name(signal)
            previous = seen_names.get(new_name)
            if previous is not None and previous != signal.name:
                raise ValueError(
                    f"Signal rename collision in frame {frame.name}: "
                    f"{signal.name} and {previous} both map to {new_name}"
                )
            seen_names[new_name] = signal.name
            renamed.append((signal, new_name))

        for signal, new_name in renamed:
            signal.name = new_name


def has_header_signals(frame: Frame) -> bool:
    return all(frame.signal_by_name(name) is not None for name in HEADER_SIGNAL_NAMES)


def is_dynamic_container_frame(frame: Frame) -> bool:
    if not has_header_signals(frame):
        return False
    if not frame.signalGroups:
        return False
    return any(signal.multiplex == "Multiplexor" for signal in frame.signals)


def is_grouped_frame_without_header(frame: Frame) -> bool:
    return bool(frame.signalGroups) and not has_header_signals(frame)


def infer_frame_size_from_signals(signals: list[Signal]) -> int:
    max_bit = 0
    for signal in signals:
        if signal.is_little_endian:
            end_bit = signal.start_bit + signal.size
        else:
            end_bit = signal.get_startbit(start_little=True) + 1
        max_bit = max(max_bit, end_bit)
    return max(1, math.ceil(max_bit / 8))


def clear_multiplex_metadata(signal: Signal) -> None:
    signal.multiplex = None
    signal.mux_val = None
    signal.mux_value = None
    signal.mux_val_grp = []
    signal.muxer_for_signal = None
    signal.is_multiplexer = False


def get_signal_mux_value(signal: Signal) -> int | None:
    if signal.multiplex not in (None, "Multiplexor"):
        return int(signal.multiplex)
    mux_val = getattr(signal, "mux_val", None)
    if mux_val is not None:
        return int(mux_val)
    mux_value = getattr(signal, "mux_value", None)
    if mux_value is not None:
        return int(mux_value)
    return None


def describe_frame_type(frame: Frame) -> str:
    if frame.is_fd:
        return "ExtendedCAN_FD" if frame.arbitration_id.extended else "StandardCAN_FD"
    return "ExtendedCAN" if frame.arbitration_id.extended else "StandardCAN"


def format_can_id_hex(can_id: int) -> str:
    return f"0x{can_id:X}"


def get_container_name(frame: Frame) -> str:
    return str(frame.attributes.get("PduName") or frame.name)


def resolve_group_pdu_name(source_frame: Frame, source_group) -> str:
    frame_pdu_name = str(source_frame.attributes.get("PduName", "")).strip()
    group_name = str(source_group.name)
    if group_name.startswith("ig") and len(group_name) > 2:
        group_name = group_name[2:]

    group_id = getattr(source_group, "id", None)
    if group_id is not None:
        suffix = f"_{group_id}"
        if group_name.endswith(suffix):
            candidate = group_name[: -len(suffix)]
            if "_PDU" in candidate:
                group_name = candidate

    if frame_pdu_name and group_name == frame_pdu_name:
        return frame_pdu_name
    if group_name:
        return group_name
    return frame_pdu_name or str(source_group.name)


BUS_CODE_MAP = {
    "ADCANFD":   0b000,  # 000
    "CHCANFD":   0b001,  # 001
    "CHCAN2":    0b010,  # 010
    "PTEXTCAN2": 0b011,  # 011
}


def bus_code_for_cluster(cluster_name: str) -> int:
    """Return the 3-bit bus code (bits 24-26) for a CAN cluster name."""
    return BUS_CODE_MAP.get(cluster_name.upper(), 0)


def make_dynamic_split_arbitration_id(
    source_frame: Frame,
    header_id: int,
    id_strategy: str,
    cluster_name: str = "",
) -> ArbitrationId:
    if id_strategy == "duplicate-container-id":
        return copy.deepcopy(source_frame.arbitration_id)
    bus_code = bus_code_for_cluster(cluster_name)
    new_id = header_id | (bus_code << 24)
    return ArbitrationId(id=new_id, extended=True)


def make_group_split_arbitration_id(
    source_frame: Frame,
    group_frame_id_mode: str,
) -> ArbitrationId:
    if group_frame_id_mode == "keep-original":
        return copy.deepcopy(source_frame.arbitration_id)
    return ArbitrationId(id=source_frame.arbitration_id.id, extended=True)


def args_like_frame_format(frame: Frame) -> str:
    return str(frame.attributes.get("VFrameFormat", ""))


def update_generated_frame_format(frame: Frame) -> None:
    if "VFrameFormat" not in frame.attributes:
        return
    if frame.is_fd:
        frame.attributes["VFrameFormat"] = (
            "ExtendedCAN_FD" if frame.arbitration_id.extended else "StandardCAN_FD"
        )
        return
    frame.attributes["VFrameFormat"] = (
        "ExtendedCAN" if frame.arbitration_id.extended else "StandardCAN"
    )


def get_dynamic_group_header_id(source_frame: Frame, source_group) -> int:
    group_signals = list(source_group.signals)
    mux_values = {
        mux_value
        for signal in group_signals
        if (mux_value := get_signal_mux_value(signal)) is not None
    }
    if not mux_values:
        raise ValueError(
            f"Signal group {source_group.name} in frame {source_frame.name} has no mux value"
        )
    if len(mux_values) != 1:
        raise ValueError(
            f"Signal group {source_group.name} in frame {source_frame.name} has multiple mux values: "
            f"{sorted(mux_values)}"
        )
    return int(next(iter(mux_values)))


def warn_skipped_dynamic_group(source_frame: Frame, source_group, reason: str) -> None:
    print(
        f"Warning: skipped dynamic signal group {source_group.name} in frame "
        f"{source_frame.name}: {reason}",
        file=sys.stderr,
    )


def build_frame_triggering_name(source_frame: Frame, pdu_name: str) -> str:
    current = str(source_frame.attributes.get("FrameTriggeringName", ""))
    if not current:
        return ""
    if pdu_name in current:
        return current
    return f"{current}__{pdu_name}"


def build_split_frame(
    cluster_name: str,
    source_frame: Frame,
    source_group,
    pdu_lengths: dict[str, int],
    shift_bits: int,
    split_kind: str,
    id_strategy: str,
    group_frame_id_mode: str,
) -> tuple[Frame, MappingRow]:
    group_signals = list(source_group.signals)
    pdu_name = resolve_group_pdu_name(source_frame, source_group)
    header_id: int | None = None

    if split_kind == "dynamic":
        header_id = get_dynamic_group_header_id(source_frame, source_group)
        arbitration_id = make_dynamic_split_arbitration_id(
            source_frame=source_frame,
            header_id=header_id,
            id_strategy=id_strategy,
            cluster_name=cluster_name,
        )
    else:
        arbitration_id = make_group_split_arbitration_id(
            source_frame=source_frame,
            group_frame_id_mode=group_frame_id_mode,
        )

    split_frame = copy.deepcopy(source_frame)
    split_frame.name = pdu_name
    split_frame.arbitration_id = arbitration_id
    split_frame.signals = []
    split_frame.signalGroups = []
    split_frame.is_complex_multiplexed = False
    split_frame.header_id = None
    split_frame.pdu_name = pdu_name
    split_frame.mux_names = {}

    split_frame.attributes["PduName"] = pdu_name
    new_triggering_name = build_frame_triggering_name(source_frame, pdu_name)
    if new_triggering_name:
        split_frame.attributes["FrameTriggeringName"] = new_triggering_name
    update_generated_frame_format(split_frame)

    for source_signal in group_signals:
        split_signal = copy.deepcopy(source_signal)
        split_signal.start_bit = source_signal.start_bit - shift_bits
        if split_signal.start_bit < 0:
            raise ValueError(
                f"Signal {split_signal.name} in frame {source_frame.name} would move below bit 0"
            )
        clear_multiplex_metadata(split_signal)
        split_frame.add_signal(split_signal)

    expected_size = pdu_lengths.get(pdu_name)
    inferred_size = infer_frame_size_from_signals(split_frame.signals)
    if expected_size is None:
        split_frame.size = inferred_size
    else:
        if inferred_size > expected_size:
            raise ValueError(
                f"PDU {pdu_name} needs at least {inferred_size} bytes from signals, "
                f"but ARXML defines {expected_size} bytes"
            )
        split_frame.size = expected_size

    mapping = MappingRow(
        source_frame_name=source_frame.name,
        source_frame_id_hex=format_can_id_hex(source_frame.arbitration_id.id),
        source_frame_type=describe_frame_type(source_frame),
        container_name=get_container_name(source_frame),
        source_group_name=str(source_group.name),
        new_msg_name=split_frame.name,
        new_msg_id_hex=format_can_id_hex(split_frame.arbitration_id.id),
        new_msg_type=describe_frame_type(split_frame),
        new_msg_dlc=split_frame.size,
        signal_count=len(split_frame.signals),
        shift_bits=shift_bits,
    )
    return split_frame, mapping


def rebuild_matrix_frames(matrix: CanMatrix, frames: list[Frame]) -> None:
    matrix.frames = []
    matrix.frames_dict_name = {}
    matrix.frames_dict_id = {}
    matrix._frames_dict_id_extend = {}
    for frame in frames:
        matrix.add_frame(frame)


def sorted_string_items(items) -> tuple[tuple[str, str], ...]:
    return tuple(sorted((str(key), str(value)) for key, value in dict(items).items()))


def signal_signature(signal: Signal) -> tuple:
    return (
        signal.name,
        signal.start_bit,
        signal.size,
        signal.is_little_endian,
        signal.is_signed,
        str(signal.factor),
        str(signal.offset),
        str(signal.min),
        str(signal.max),
        signal.unit,
        tuple(sorted(str(receiver) for receiver in signal.receivers)),
        tuple(sorted((str(key), str(value)) for key, value in signal.values.items())),
        sorted_string_items(signal.attributes),
    )


def frame_signature(frame: Frame) -> tuple:
    comparable_attributes = {
        key: value
        for key, value in frame.attributes.items()
        if key != "FrameTriggeringName"
    }
    return (
        frame.name,
        frame.arbitration_id.id,
        frame.arbitration_id.extended,
        frame.size,
        frame.is_fd,
        tuple(sorted(str(transmitter) for transmitter in frame.transmitters)),
        sorted_string_items(comparable_attributes),
        tuple(signal_signature(signal) for signal in frame.signals),
    )


def describe_frame_origin(frame: Frame) -> str:
    triggering_name = frame.attributes.get("FrameTriggeringName")
    if triggering_name:
        return f"{frame.name} ({triggering_name})"
    return frame.name


def deduplicate_identical_frame_ids(frames: list[Frame]) -> tuple[list[Frame], int]:
    unique_frames: list[Frame] = []
    seen: dict[tuple[int, bool], tuple[Frame, tuple]] = {}
    duplicate_count = 0

    for frame in frames:
        key = (frame.arbitration_id.id, frame.arbitration_id.extended)
        signature = frame_signature(frame)
        previous = seen.get(key)
        if previous is None:
            seen[key] = (frame, signature)
            unique_frames.append(frame)
            continue

        previous_frame, previous_signature = previous
        if signature == previous_signature:
            duplicate_count += 1
            continue

        frame_type = "extended" if frame.arbitration_id.extended else "standard"
        raise ValueError(
            "Duplicate frame ID detected after split with different frame definitions: "
            f"{describe_frame_origin(frame)} and {describe_frame_origin(previous_frame)} "
            f"both use ID {frame.arbitration_id.id} ({frame_type})"
        )

    return unique_frames, duplicate_count


def get_dynamic_shift_bits(frame: Frame) -> int:
    header_signals = [frame.signal_by_name(name) for name in HEADER_SIGNAL_NAMES]
    if any(signal is None for signal in header_signals):
        raise ValueError(f"Dynamic frame {frame.name} is missing one of {HEADER_SIGNAL_NAMES}")
    return max(signal.start_bit + signal.size for signal in header_signals if signal is not None)


def split_frames(
    cluster_name: str,
    matrix: CanMatrix,
    pdu_lengths: dict[str, int],
    id_strategy: str,
    group_frame_id_mode: str,
) -> tuple[CanMatrix, SplitStats, list[MappingRow]]:
    output_matrix = copy.deepcopy(matrix)
    rebuilt_frames: list[Frame] = []
    stats = SplitStats()
    mapping_rows: list[MappingRow] = []

    for frame in output_matrix.frames:
        if is_dynamic_container_frame(frame):
            stats.dynamic_frame_count += 1
            shift_bits = get_dynamic_shift_bits(frame)

            for group in frame.signalGroups:
                try:
                    split_frame, mapping = build_split_frame(
                        cluster_name=cluster_name,
                        source_frame=frame,
                        source_group=group,
                        pdu_lengths=pdu_lengths,
                        shift_bits=shift_bits,
                        split_kind="dynamic",
                        id_strategy=id_strategy,
                        group_frame_id_mode=group_frame_id_mode,
                    )
                except ValueError as exc:
                    if "has no mux value" not in str(exc):
                        raise
                    warn_skipped_dynamic_group(frame, group, "group has no mux value")
                    stats.skipped_dynamic_group_count += 1
                    continue
                rebuilt_frames.append(split_frame)
                mapping_rows.append(mapping)
                stats.created_split_count += 1
            continue

        if is_grouped_frame_without_header(frame):
            stats.preserved_grouped_frame_count += 1
            rebuilt_frames.append(frame)
            continue

        if not is_dynamic_container_frame(frame):
            rebuilt_frames.append(frame)

    rebuilt_frames, stats.duplicate_split_count = deduplicate_identical_frame_ids(rebuilt_frames)
    stats.created_split_count -= stats.duplicate_split_count
    rebuild_matrix_frames(output_matrix, rebuilt_frames)
    return output_matrix, stats, mapping_rows


def ensure_unique_ids(matrix: CanMatrix) -> None:
    seen: dict[tuple[int, bool], str] = {}
    for frame in matrix.frames:
        key = (frame.arbitration_id.id, frame.arbitration_id.extended)
        previous = seen.get(key)
        if previous is not None:
            raise ValueError(
                f"Duplicate frame ID detected after split: {frame.name} and {previous} "
                f"both use ID {frame.arbitration_id.id} "
                f"({'extended' if frame.arbitration_id.extended else 'standard'})"
            )
        seen[key] = frame.name


def dump_can_matrix(
    matrix: CanMatrix,
    output_path: Path,
    dbc_export_encoding: str,
    dbc_export_comment_encoding: str,
    zero_factor_policy: str,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    validate_exportable_signals(matrix, zero_factor_policy)
    canmatrix.formats.dumpp(
        {"": matrix},
        str(output_path),
        export_type="dbc",
        dbcExportEncoding=dbc_export_encoding,
        dbcExportCommentEncoding=dbc_export_comment_encoding,
    )


def iter_export_validation_issues(matrix: CanMatrix) -> list[ExportValidationIssue]:
    issues: list[ExportValidationIssue] = []
    for frame in matrix.frames:
        for signal in frame.signals:
            try:
                signal.phys2raw(None)
            except Exception as exc:
                issues.append(
                    ExportValidationIssue(
                        frame_name=frame.name,
                        signal_name=signal.name,
                        factor=str(signal.factor),
                        offset=str(signal.offset),
                        minimum=str(signal.min),
                        maximum=str(signal.max),
                        error_type=type(exc).__name__,
                        error_message=str(exc),
                    )
                )
    return issues


def format_export_validation_issues(issues: list[ExportValidationIssue], limit: int = 20) -> str:
    lines = [
        "DBC export failed because some signals cannot be converted by canmatrix.",
        f"Problematic signals: {len(issues)}",
    ]
    for issue in issues[:limit]:
        lines.append(
            f"- Frame={issue.frame_name}, Signal={issue.signal_name}, "
            f"factor={issue.factor}, offset={issue.offset}, min={issue.minimum}, "
            f"max={issue.maximum}, error={issue.error_type}: {issue.error_message}"
        )
    remaining = len(issues) - min(len(issues), limit)
    if remaining > 0:
        lines.append(f"... and {remaining} more")
    lines.append(
        "Fix the source ARXML/DBC scaling, or rerun with "
        "--zero-factor-policy set-one to replace factor 0 with 1 during export."
    )
    return "\n".join(lines)


def replace_zero_factors(matrix: CanMatrix) -> int:
    changed = 0
    for frame in matrix.frames:
        for signal in frame.signals:
            if signal.factor == 0:
                signal.factor = 1
                changed += 1
    return changed


def validate_exportable_signals(matrix: CanMatrix, zero_factor_policy: str) -> None:
    if zero_factor_policy == "set-one":
        changed = replace_zero_factors(matrix)
        if changed:
            print(
                f"Warning: replaced factor=0 with factor=1 for {changed} signal(s) before DBC export.",
                file=sys.stderr,
            )

    issues = iter_export_validation_issues(matrix)
    if issues:
        raise ValueError(format_export_validation_issues(issues))


def autofit_worksheet_columns(worksheet) -> None:
    for column_index, column_cells in enumerate(worksheet.columns, start=1):
        cell_lengths = [len(str(cell.value)) for cell in column_cells if cell.value is not None]
        if not cell_lengths:
            continue
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(cell_lengths) + 2,
            48,
        )


def estimate_wrapped_line_count(text: str, width_chars: float) -> int:
    usable_width = max(int(width_chars), 1)
    total_lines = 0
    for raw_line in str(text).splitlines() or [""]:
        line = raw_line or " "
        total_lines += max(1, math.ceil(len(line) / usable_width))
    return max(total_lines, 1)


def autofit_worksheet_rows(worksheet) -> None:
    default_height = 18
    merged_lookup: dict[tuple[int, int], tuple[int, int]] = {}
    for merged_range in worksheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                merged_lookup[(row, col)] = (max_row - min_row + 1, max_col - min_col + 1)

    for row_index in range(1, worksheet.max_row + 1):
        max_lines = 1
        for col_index in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row_index, col_index)
            if cell.value is None:
                continue

            column_letter = cell.column_letter
            width_chars = worksheet.column_dimensions[column_letter].width or 10
            row_span, col_span = merged_lookup.get((row_index, col_index), (1, 1))
            width_chars = max(width_chars * col_span - 1, 1)
            lines = estimate_wrapped_line_count(str(cell.value), width_chars)
            max_lines = max(max_lines, math.ceil(lines / row_span))

        worksheet.row_dimensions[row_index].height = max(default_height, max_lines * default_height)


def merge_repeated_source_cells(worksheet) -> None:
    if worksheet.max_row <= 2:
        return

    start_row = 2
    merge_columns = range(1, 5)
    current_key = tuple(worksheet.cell(start_row, column).value for column in merge_columns)

    for row in range(3, worksheet.max_row + 2):
        next_key = None
        if row <= worksheet.max_row:
            next_key = tuple(worksheet.cell(row, column).value for column in merge_columns)
        if next_key == current_key:
            continue
        if row - start_row > 1:
            for column in merge_columns:
                worksheet.merge_cells(
                    start_row=start_row,
                    start_column=column,
                    end_row=row - 1,
                    end_column=column,
                )
                worksheet.cell(start_row, column).alignment = Alignment(
                    vertical="center",
                    horizontal="center",
                    wrap_text=True,
                )
        start_row = row
        current_key = next_key


def style_summary_sheet(worksheet) -> None:
    worksheet.insert_rows(1)
    worksheet.merge_cells("A1:B1")
    title_cell = worksheet["A1"]
    title_cell.value = "Mapping Summary"
    title_cell.font = Font(name="Arial", bold=True, size=13, color="FFFFFF")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.fill = PatternFill(fill_type="solid", fgColor="4F81BD")
    worksheet.row_dimensions[1].height = 24


def style_mapping_cells(worksheet, border: Border) -> None:
    left_align_columns = {1, 4, 5, 6}
    center_align_columns = {2, 3, 7, 8, 9, 10, 11}

    for row_index, row in enumerate(worksheet.iter_rows(), start=1):
        for cell in row:
            cell.border = border
            if row_index == 1:
                continue
            if cell.column in left_align_columns:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif cell.column in center_align_columns:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)


def style_summary_cells(worksheet, border: Border) -> None:
    for row_index, row in enumerate(worksheet.iter_rows(), start=1):
        for cell in row:
            cell.border = border
            if row_index == 1:
                continue
            if cell.column == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def dump_mapping_workbook(
    mapping_rows: list[MappingRow],
    stats: SplitStats,
    output_path: Path,
    cluster_name: str,
    source_label: str,
    output_dbc_name: str,
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Item", "Value"])
    for row in [
        ("Cluster", cluster_name),
        ("Source", source_label),
        ("OutputDbc", output_dbc_name),
        ("DynamicSourceFrames", stats.dynamic_frame_count),
        ("PreservedGroupedFramesWithoutHeader", stats.preserved_grouped_frame_count),
        ("GeneratedMessages", stats.created_split_count),
        ("DeduplicatedMessages", stats.duplicate_split_count),
        ("MappingRows", len(mapping_rows)),
    ]:
        summary.append(list(row))
    style_summary_sheet(summary)

    mappings = workbook.create_sheet("Mappings")
    mappings.append(MAPPING_HEADERS)
    for item in mapping_rows:
        mappings.append(item.to_row())
    merge_repeated_source_cells(mappings)

    header_font = Font(name="Arial", bold=True)
    body_font = Font(name="Arial")
    header_fill = PatternFill(fill_type="solid", fgColor="9CC2E5")
    thin_side = Side(style="thin", color="A6A6A6")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        for row_index, row in enumerate(worksheet.iter_rows(), start=1):
            for cell in row:
                if worksheet.title == "Summary" and row_index == 1:
                    continue
                cell.font = header_font if row_index == 1 else body_font
                if row_index == 1:
                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical="center",
                        wrap_text=True,
                    )
                    cell.fill = header_fill
        if worksheet.title == "Mappings":
            style_mapping_cells(worksheet, border)
        else:
            style_summary_cells(worksheet, border)
        if worksheet.max_column > 0 and worksheet.max_row > 0:
            worksheet.auto_filter.ref = worksheet.dimensions
        autofit_worksheet_columns(worksheet)
        autofit_worksheet_rows(worksheet)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def derive_output_prefix(arxml_path: Path, output_prefix: str | None) -> str:
    if output_prefix:
        return output_prefix
    return arxml_path.stem


def build_raw_output_path(output_dir: Path, output_prefix: str, cluster_name: str) -> Path:
    return output_dir / f"{output_prefix}_{cluster_name}.dbc"


def build_split_output_path(
    output_dir: Path,
    output_prefix: str,
    cluster_name: str,
    suffix: str,
) -> Path:
    return output_dir / f"{output_prefix}_{cluster_name}{suffix}.dbc"


def build_mapping_output_path(output_path: Path) -> Path:
    return output_path.with_name(f"{output_path.stem}_mapping.xlsx")


def format_deduplicated_suffix(stats: SplitStats) -> str:
    parts: list[str] = []
    if stats.duplicate_split_count:
        parts.append(f"deduplicated {stats.duplicate_split_count} identical messages")
    if stats.skipped_dynamic_group_count:
        parts.append(f"skipped {stats.skipped_dynamic_group_count} dynamic groups without mux value")
    if stats.preserved_grouped_frame_count:
        parts.append(
            f"preserved {stats.preserved_grouped_frame_count} grouped frames without Header_ID/Header_DLC"
        )
    if not parts:
        return ""
    return ", " + ", ".join(parts)


def process_one_matrix(
    cluster_name: str,
    matrix: CanMatrix,
    pdu_lengths: dict[str, int],
    output_path: Path,
    id_strategy: str,
    group_frame_id_mode: str,
    mapping_xlsx: bool,
    source_label: str,
    dbc_export_encoding: str,
    dbc_export_comment_encoding: str,
    zero_factor_policy: str,
) -> tuple[SplitStats, Path | None]:
    split_matrix, stats, mapping_rows = split_frames(
        cluster_name=cluster_name,
        matrix=matrix,
        pdu_lengths=pdu_lengths,
        id_strategy=id_strategy,
        group_frame_id_mode=group_frame_id_mode,
    )

    if id_strategy != "duplicate-container-id":
        ensure_unique_ids(split_matrix)

    dump_can_matrix(
        split_matrix,
        output_path,
        dbc_export_encoding=dbc_export_encoding,
        dbc_export_comment_encoding=dbc_export_comment_encoding,
        zero_factor_policy=zero_factor_policy,
    )

    mapping_path: Path | None = None
    if mapping_xlsx:
        mapping_path = build_mapping_output_path(output_path)
        dump_mapping_workbook(
            mapping_rows=mapping_rows,
            stats=stats,
            output_path=mapping_path,
            cluster_name=cluster_name,
            source_label=source_label,
            output_dbc_name=output_path.name,
        )

    return stats, mapping_path


def run_split_only(args: argparse.Namespace, pdu_lengths: dict[str, int]) -> int:
    output_dir = Path(args.output_dir).resolve()
    dbc_files = expand_dbc_inputs(args.inputs)

    total_dynamic_frames = 0
    total_preserved_grouped_frames = 0
    total_split_frames = 0
    total_duplicate_frames = 0
    for dbc_path in dbc_files:
        cluster_name, matrix = load_can_matrix(dbc_path)
        output_path = output_dir / f"{dbc_path.stem}{args.suffix}{dbc_path.suffix}"
        stats, mapping_path = process_one_matrix(
            cluster_name=cluster_name,
            matrix=matrix,
            pdu_lengths=pdu_lengths,
            output_path=output_path,
            id_strategy=args.id_strategy,
            group_frame_id_mode=args.group_frame_id_mode,
            mapping_xlsx=args.mapping_xlsx,
            source_label=dbc_path.name,
            dbc_export_encoding=args.dbc_export_encoding,
            dbc_export_comment_encoding=args.dbc_export_comment_encoding,
            zero_factor_policy=args.zero_factor_policy,
        )

        total_dynamic_frames += stats.dynamic_frame_count
        total_preserved_grouped_frames += stats.preserved_grouped_frame_count
        total_split_frames += stats.created_split_count
        total_duplicate_frames += stats.duplicate_split_count
        mapping_suffix = f", {mapping_path.name} created" if mapping_path else ""
        print(
            f"{dbc_path.name} -> {output_path.name}: "
            f"split {stats.dynamic_frame_count} dynamic frames into "
            f"{stats.created_split_count} messages"
            f"{format_deduplicated_suffix(stats)}{mapping_suffix}"
        )

    print(
        f"Done. Processed {len(dbc_files)} DBC file(s), "
        f"split {total_dynamic_frames} dynamic frame(s) into {total_split_frames} message(s)"
        f"{', preserved ' + str(total_preserved_grouped_frames) + ' grouped frame(s) without Header_ID/Header_DLC' if total_preserved_grouped_frames else ''}"
        f"{format_total_deduplicated_suffix(total_duplicate_frames)}."
    )
    return 0


def format_total_deduplicated_suffix(duplicate_count: int) -> str:
    if not duplicate_count:
        return ""
    return f", deduplicated {duplicate_count} identical message(s)"


def run_convert_and_split(args: argparse.Namespace, arxml_path: Path, pdu_lengths: dict[str, int]) -> int:
    output_dir = Path(args.output_dir).resolve()
    output_prefix = derive_output_prefix(arxml_path, args.output_prefix)
    can_cluster = load_can_cluster_from_arxml(arxml_path)

    total_dynamic_frames = 0
    total_preserved_grouped_frames = 0
    total_split_frames = 0
    total_duplicate_frames = 0
    for cluster_name, matrix in can_cluster.items():
        raw_output_path = build_raw_output_path(output_dir, output_prefix, cluster_name)
        split_output_path = build_split_output_path(output_dir, output_prefix, cluster_name, args.suffix)

        if args.keep_converted_dbc:
            dump_can_matrix(
                matrix,
                raw_output_path,
                dbc_export_encoding=args.dbc_export_encoding,
                dbc_export_comment_encoding=args.dbc_export_comment_encoding,
                zero_factor_policy=args.zero_factor_policy,
            )

        stats, mapping_path = process_one_matrix(
            cluster_name=cluster_name,
            matrix=matrix,
            pdu_lengths=pdu_lengths,
            output_path=split_output_path,
            id_strategy=args.id_strategy,
            group_frame_id_mode=args.group_frame_id_mode,
            mapping_xlsx=args.mapping_xlsx,
            source_label=arxml_path.name,
            dbc_export_encoding=args.dbc_export_encoding,
            dbc_export_comment_encoding=args.dbc_export_comment_encoding,
            zero_factor_policy=args.zero_factor_policy,
        )

        total_dynamic_frames += stats.dynamic_frame_count
        total_preserved_grouped_frames += stats.preserved_grouped_frame_count
        total_split_frames += stats.created_split_count
        total_duplicate_frames += stats.duplicate_split_count

        mapping_suffix = f", {mapping_path.name} created" if mapping_path else ""
        if args.keep_converted_dbc:
            print(
                f"{cluster_name}: {raw_output_path.name} created, "
                f"{split_output_path.name} created "
                f"(split {stats.dynamic_frame_count} dynamic frames into "
                f"{stats.created_split_count} messages"
                f"{format_deduplicated_suffix(stats)}{mapping_suffix})"
            )
        else:
            print(
                f"{cluster_name}: {split_output_path.name} created "
                f"(split {stats.dynamic_frame_count} dynamic frames into "
                f"{stats.created_split_count} messages"
                f"{format_deduplicated_suffix(stats)}{mapping_suffix})"
            )

    print(
        f"Done. Converted {len(can_cluster)} cluster(s) from ARXML and "
        f"split {total_dynamic_frames} dynamic frame(s) into {total_split_frames} message(s)"
        f"{', preserved ' + str(total_preserved_grouped_frames) + ' grouped frame(s) without Header_ID/Header_DLC' if total_preserved_grouped_frames else ''}"
        f"{format_total_deduplicated_suffix(total_duplicate_frames)}."
    )
    return 0


def main() -> int:
    args = parse_args()
    mode = validate_args(args)
    warn_if_risky_options(args)
    arxml_path = Path(args.arxml).resolve()
    pdu_lengths = parse_pdu_lengths(arxml_path)

    if mode == "split-only":
        return run_split_only(args, pdu_lengths)
    return run_convert_and_split(args, arxml_path, pdu_lengths)


if __name__ == "__main__":
    sys.exit(main())
